#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
命令行模式：在无图形界面环境下运行 Excel 工具的核心功能

示例：
  列对比：
    python cli.py compare --file-a sample_file_a.xlsx --col-a 姓名 \
                         --file-b sample_file_b.xlsx --col-b 姓名 \
                         --output compare_result.xlsx

  数据去重：
    python cli.py dedup --file sample_file_a.xlsx --column 姓名 --keep first \
                        --output dedup_result.xlsx

  格式调整（日期/数字/文本）：
    python cli.py format --file sample_file_a.xlsx --column 入职日期 --type date \
                         --output format_result.xlsx
"""

import argparse
import sys
import pandas as pd
from openpyxl.styles import PatternFill


def compare_columns(file_a: str, col_a: str, file_b: str, col_b: str, output: str) -> None:
    df_a = pd.read_excel(file_a)
    df_b = pd.read_excel(file_b)

    data_a = set(df_a[col_a].dropna().astype(str))
    data_b = set(df_b[col_b].dropna().astype(str))

    only_in_a = data_a - data_b
    only_in_b = data_b - data_a

    result_df = df_a.copy()
    mask_a = result_df[col_a].astype(str).isin(only_in_a)
    result_df['差异标记'] = ''
    result_df.loc[mask_a, '差异标记'] = 'A中独有'

    b_only_data = df_b[df_b[col_b].astype(str).isin(only_in_b)]
    if not b_only_data.empty:
        new_rows = []
        for _, row in b_only_data.iterrows():
            new_row = {c: None for c in result_df.columns}
            new_row[col_a] = row[col_b]
            new_row['差异标记'] = 'B中独有'
            new_rows.append(new_row)
        if new_rows:
            result_df = pd.concat([result_df, pd.DataFrame(new_rows)], ignore_index=True)

    # 保存并标红差异行
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        result_df.to_excel(writer, index=False, sheet_name='对比结果')
        ws = writer.sheets['对比结果']
        red_fill = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')
        for row in ws.iter_rows(min_row=2):
            if row[-1].value in ['A中独有', 'B中独有']:
                for cell in row:
                    cell.fill = red_fill

    print(f"对比完成：仅A有 {len(only_in_a)} 条，仅B有 {len(only_in_b)} 条；已保存 {output}")


def dedup_file(file_path: str, column: str, keep: str, output: str) -> None:
    df = pd.read_excel(file_path)
    if keep not in {"first", "last"}:
        raise ValueError("keep 只能为 'first' 或 'last'")
    original = len(df)
    result = df.drop_duplicates(subset=[column], keep=keep)
    removed = original - len(result)
    result.to_excel(output, index=False)
    print(f"去重完成：原始 {original} 条，删除 {removed} 条，结果 {len(result)} 条；已保存 {output}")


def format_column(file_path: str, column: str, fmt_type: str, output: str) -> None:
    df = pd.read_excel(file_path)
    result = df.copy()

    if fmt_type == 'date':
        result[column] = pd.to_datetime(result[column], errors='coerce').dt.strftime('%Y-%m-%d')
    elif fmt_type == 'number':
        result[column] = pd.to_numeric(result[column], errors='coerce').round(2)
    elif fmt_type == 'text':
        result[column] = result[column].astype(str).str.strip()
    else:
        raise ValueError("type 只能为 'date' | 'number' | 'text'")

    result.to_excel(output, index=False)
    print(f"格式调整完成：类型 {fmt_type}，已保存 {output}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description='Excel 工具（命令行模式）')
    sub = parser.add_subparsers(dest='cmd', required=True)

    p_cmp = sub.add_parser('compare', help='列对比')
    p_cmp.add_argument('--file-a', required=True)
    p_cmp.add_argument('--col-a', required=True)
    p_cmp.add_argument('--file-b', required=True)
    p_cmp.add_argument('--col-b', required=True)
    p_cmp.add_argument('--output', required=True)

    p_dedup = sub.add_parser('dedup', help='数据去重')
    p_dedup.add_argument('--file', required=True)
    p_dedup.add_argument('--column', required=True)
    p_dedup.add_argument('--keep', choices=['first', 'last'], default='first')
    p_dedup.add_argument('--output', required=True)

    p_fmt = sub.add_parser('format', help='格式调整')
    p_fmt.add_argument('--file', required=True)
    p_fmt.add_argument('--column', required=True)
    p_fmt.add_argument('--type', choices=['date', 'number', 'text'], required=True)
    p_fmt.add_argument('--output', required=True)

    return parser


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        if args.cmd == 'compare':
            compare_columns(args.file_a, args.col_a, args.file_b, args.col_b, args.output)
        elif args.cmd == 'dedup':
            dedup_file(args.file, args.column, args.keep, args.output)
        elif args.cmd == 'format':
            format_column(args.file, args.column, args.type, args.output)
    except Exception as exc:
        print(f"错误：{exc}")
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())